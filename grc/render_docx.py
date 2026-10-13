"""Render the analyst orientation brief (.docx) from engagement data.

Layout and styles come from templates/brief-template.docx so the output
matches the house style; all content comes from the repo.
"""
from __future__ import annotations

import copy
import re
import datetime as dt
import zipfile
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .model import LEGS, Engagement
from .rules import metrics

NAVY = "183B4E"
CALLOUT = "EAF2F3"
BAND = "F3F3F3"
WHITE = "FFFFFF"
FULL = 10080        # callout width (dxa)
CONTENT = 10368     # full table width (dxa)
BULLET_NUM_ID = "1"  # bullet list defined in the template's numbering.xml


# ------------------------------------------------------------------ helpers
def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _run(p, text: str, bold=False, color: str | None = None, size: float | None = None):
    r = p.add_run(_CTRL.sub("\ufffd", str(text)))   # Word XML can't hold control characters
    r.bold = bold or None
    if color:
        rpr = r._r.get_or_add_rPr()
        rpr.append(_el("w:color", **{"w:val": color}))
    if size:
        r.font.size = Pt(size)
    return r


def _spacing(p, after: int | None = None):
    ppr = p._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = _el("w:spacing")
        ppr.append(sp)
    if after is not None:
        sp.set(qn("w:after"), str(after))


def bullet(doc, text: str, last=False):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    num = _el("w:numPr")
    num.append(_el("w:ilvl", **{"w:val": 0}))
    num.append(_el("w:numId", **{"w:val": BULLET_NUM_ID}))
    ppr.append(num)
    ppr.append(_el("w:spacing", **{"w:after": 40 if last else 0}))
    ppr.append(_el("w:ind", **{"w:left": 259}))
    _run(p, text)
    return p


def bullets(doc, items):
    items = list(items)
    for i, t in enumerate(items):
        bullet(doc, t, last=i == len(items) - 1)


def heading(doc, text: str, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def _table(doc, widths: list[int], style_id: str):
    t = doc.add_table(rows=0, cols=len(widths))
    tbl = t._tbl
    tblPr = tbl.tblPr
    for child in list(tblPr):
        tblPr.remove(child)
    tblPr.append(_el("w:tblStyle", **{"w:val": style_id}))
    tblPr.append(_el("w:tblW", **{"w:w": sum(widths), "w:type": "dxa"}))
    tblPr.append(_el("w:jc", **{"w:val": "center"}))
    tblPr.append(_el("w:tblLayout", **{"w:type": "fixed"}))
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    return t


def _cell_fmt(cell, width: int, fill: str | None, pad: int | None = None, vcenter=False):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(_el("w:tcW", **{"w:w": width, "w:type": "dxa"}))
    if fill:
        tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
    if pad:
        mar = _el("w:tcMar")
        mar.append(_el("w:top", **{"w:w": pad, "w:type": "dxa"}))
        mar.append(_el("w:bottom", **{"w:w": pad, "w:type": "dxa"}))
        tcPr.append(mar)
    if vcenter:
        tcPr.append(_el("w:vAlign", **{"w:val": "center"}))


def _row(t, header=False):
    row = t.add_row()
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:cantSplit"))
    if header:
        trPr.append(_el("w:tblHeader"))
    trPr.append(_el("w:jc", **{"w:val": "center"}))
    return row


def _keep_with_next(cell):
    for p in cell.paragraphs:
        p.paragraph_format.keep_with_next = True


def callout(doc, text: str):
    t = _table(doc, [FULL], "a")
    row = _row(t)
    cell = row.cells[0]
    _cell_fmt(cell, FULL, CALLOUT, pad=120, vcenter=True)
    p = cell.paragraphs[0]
    p.alignment = 1
    _run(p, text, bold=True, color=NAVY)
    spacer = doc.add_paragraph()
    _spacing(spacer, 20)


def data_table(doc, headers: list[str], rows: list[list], widths: list[int], first_col_color: str | None = None,
               bulleted_cols: tuple[int, ...] = (), bold_cols: tuple[int, ...] = (0,), pad=75):
    """rows: list of cell values; a list value in a bulleted column becomes '• ' lines."""
    t = _table(doc, widths, "a0")
    hr = _row(t, header=True)
    for i, h in enumerate(headers):
        c = hr.cells[i]
        _cell_fmt(c, widths[i], NAVY, pad=90 if bulleted_cols else None)
        _run(c.paragraphs[0], h.upper(), bold=True, color=WHITE)
        _keep_with_next(c)
    for ri, values in enumerate(rows):
        r = _row(t)
        fill = BAND if ri % 2 == 0 else None
        for ci, v in enumerate(values):
            c = r.cells[ci]
            _cell_fmt(c, widths[ci], fill, pad=pad)
            if ci in bulleted_cols and isinstance(v, list):
                p0 = c.paragraphs[0]
                for li, line in enumerate(v):
                    p = p0 if li == 0 else c.add_paragraph()
                    ppr = p._p.get_or_add_pPr()
                    ppr.append(_el("w:spacing", **{"w:after": 20}))
                    ppr.append(_el("w:ind", **{"w:left": 115}))
                    _run(p, f"• {line}")
            else:
                color = first_col_color if ci == 0 else None
                _run(c.paragraphs[0], str(v), bold=ci in bold_cols, color=color)
    doc.add_paragraph()
    return t


# -------------------------------------------------------------------- brief
def render_brief(model: Engagement, out_path: Path) -> Path:
    from .model import resource
    template = resource(model.root, "templates", "brief-template.docx")
    doc = docx.Document(str(template))
    body = doc.element.body
    for p in list(body):  # drop the placeholder paragraph, keep sectPr
        if p.tag != qn("w:sectPr"):
            body.remove(p)

    e = model.engagement
    fw = model.catalog["framework"]
    pb = model.playbook
    client = model.client

    # Title block
    title = doc.add_paragraph(style="Title")
    title.alignment = 1
    title.add_run(f"{client} {pb.get('title_suffix', fw['short_name'] + ' Project Brief')}")
    sub = doc.add_paragraph()
    _spacing(sub, 200)
    _run(sub, e.get("subtitle", "Analyst Orientation"), bold=True, color=NAVY, size=11)
    doc.add_paragraph()
    callout(doc, "TLDR: " + model.fmt(e.get("tldr") or pb.get("tldr_template", "")))

    # 1. Client overview
    heading(doc, "1. Client Overview")
    bullets(doc, [model.fmt(x) for x in e["client"].get("overview", [])])

    # 2. AI usage
    heading(doc, "2. AI Usage")
    bullets(doc, [model.fmt(x) for x in e.get("ai_usage", [])])

    # 3. In place
    heading(doc, "3. What is already in place")
    doc.add_paragraph()
    data_table(doc, ["Foundation", "Examples"],
               [[x["foundation"], model.fmt(x["examples"])] for x in e.get("in_place", [])],
               [5184, 5184])

    # 4. Gap analysis
    heading(doc, "4. Gap Analysis")
    callout(doc, model.fmt(pb.get("gap_callout", "")))
    gap_rows = [[g["area"], model.fmt(g["finding"])] for g in model.gaps
                if g.get("in_brief", True) and g.get("status") in ("open", "in_progress")]
    data_table(doc, ["Area", "What to verify / close"], gap_rows, [3305, 7075])

    heading(doc, "Analyst evidence test", 2)
    bullets(doc, [f"{t['label']} — {model.fmt(t['question'])}" for t in pb["evidence_test"]])

    m = metrics(model)
    heading(doc, f"Readiness snapshot (as of {m['as_of']})", 2)
    p = doc.add_paragraph()
    _run(p, f"{m['requirements_ready']} of {m['requirements_applicable']} applicable requirements pass all four tests; "
            f"{m['gaps_open']} gaps open; {m['legs_verified_pct']}% of evidence-test legs verified. "
            f"Generated from the engagement repository - do not edit by hand.")
    p.paragraph_format.keep_with_next = True
    labels = {t["key"]: t["label"] for t in pb["evidence_test"]}
    data_table(doc, ["Test", "Verified", "Partial", "Claimed", "Not started"],
               [[labels[k], m["legs"][k]["verified"], m["legs"][k]["partial"], m["legs"][k]["claimed"],
                 m["legs"][k]["not_started"]] for k in LEGS],
               [3168, 1800, 1800, 1800, 1800], first_col_color=None)

    # 5. Timeline
    heading(doc, "5. Timeline")
    data_table(doc, ["Month", "Objective", "Deliverables"],
               [[ph["label"], ph["objective"], [model.fmt(d) for d in ph["deliverables"]]] for ph in model.phases],
               [1656, 2232, 6408], first_col_color=NAVY, bulleted_cols=(2,), bold_cols=(0, 1), pad=80)

    # 6. Needs
    heading(doc, "6. Needs")
    bullets(doc, [model.fmt(x) for x in model.setting("needs", [])])
    heading(doc, "Working rule", 2)
    callout(doc, model.fmt(model.setting("working_rule", "")))

    # 7. Requirements
    heading(doc, f"7. {pb.get('requirements_heading', 'ISO Requirements')}")
    bullets(doc, [model.fmt(x) for x in pb.get("readiness_criteria", [])])

    # 8. Guardrails
    heading(doc, "8. Engagement Guardrails")
    b = e["budget"]
    from .hours import hours_used, tracked
    g = [f"{b['hours_total']:g} total consulting hours" + (f" ({hours_used(model):g} used)" if tracked(model) else "")
         + (f": {b['note']}" if b.get("note") else ".")]
    g += [model.fmt(x) for x in model.setting("guardrails", pb.get("default_guardrails", []))]
    bullets(doc, g)

    # 9. Risks
    heading(doc, "9. Key Risks")
    bullets(doc, [model.fmt(x) for x in model.setting("key_risks", pb.get("default_risks", []))])

    # metadata
    cp = doc.core_properties
    cp.title = f"{client} {fw['short_name']} Project Brief"
    cp.author = "grc-readiness"
    cp.last_modified_by = "grc-readiness"

    stamp = dt.datetime.combine(model.today, dt.time(12, 0))
    cp.created = cp.modified = stamp
    cp.revision = 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    _normalize_zip(out_path, stamp)
    return out_path


def _normalize_zip(path: Path, stamp: dt.datetime) -> None:
    """Fixed zip timestamps so the same data always produces the same file (clean diffs, stable hashes)."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=stamp.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zout.writestr(info, zin.read(item.filename))
    tmp.replace(path)


# ---------------------------------------------------------------- policies and procedures
def _inline(p, text: str, bold=False, color=None):
    """**bold** and `code` (rendered plain) inside a paragraph."""
    for i, part in enumerate(text.replace("`", "").split("**")):
        if part:
            _run(p, part, bold=bold or i % 2 == 1, color=color)


def _md_blocks(body: str):
    """Split the Markdown subset used by the templates into blocks."""
    lines = body.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            yield ("h", min(level, 3), s.lstrip("#").strip())
            i += 1
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            yield ("table", rows)
        elif s[:2] in ("- ", "* "):
            yield ("bullet", s[2:])
            i += 1
        elif s[0].isdigit() and ". " in s[:5]:
            yield ("num", s)
            i += 1
        else:
            para = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#", "|", "- ", "* ")) \
                    and not (lines[i].strip()[0].isdigit() and ". " in lines[i].strip()[:5]):
                para.append(lines[i].strip())
                i += 1
            yield ("p", " ".join(para))


def render_markdown(model: Engagement, body: str, out_path: Path, title: str, stamp_date: dt.date,
                    banner: str | None = None, control: list[list[str]] | None = None, control_title: str = "Document control") -> Path:
    """Markdown subset -> Word in the house style (shared by policies, procedures and review packs)."""
    from .model import resource
    doc = docx.Document(str(resource(model.root, "templates", "brief-template.docx")))
    root = doc.element.body
    for p in list(root):
        if p.tag != qn("w:sectPr"):
            root.remove(p)
    blocks = list(_md_blocks(body))
    if blocks and blocks[0][0] == "h" and blocks[0][1] == 1:
        title = blocks.pop(0)[2]
    t = doc.add_paragraph(style="Title")
    _run(t, title)
    if banner:
        callout(doc, banner)
    if control:
        data_table(doc, [control_title, ""], control, [2600, CONTENT - 2600])
    for b in blocks:
        if b[0] == "h":
            heading(doc, b[2], level=max(1, b[1] - 1))
        elif b[0] == "p":
            p = doc.add_paragraph()
            _inline(p, b[1])
        elif b[0] == "bullet":
            p = bullet(doc, "")
            _inline(p, b[1])
        elif b[0] == "num":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            _inline(p, b[1])
        elif b[0] == "table" and b[1]:
            header, rows = b[1][0], b[1][1:]
            n = len(header)
            rows = [(r + [""] * n)[:n] for r in rows]
            data_table(doc, [h.replace("**", "") for h in header], [[c.replace("**", "") for c in r] for r in rows],
                       [CONTENT // n] * n, bold_cols=())
    cp = doc.core_properties
    cp.title = title
    cp.author = cp.last_modified_by = "grc-readiness"
    stamp = dt.datetime.combine(stamp_date, dt.time(12, 0))
    cp.created = cp.modified = stamp
    cp.revision = 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    _normalize_zip(out_path, stamp)
    return out_path


def render_document(model: Engagement, doc_id: str, out_path: Path) -> Path:
    """One policy/procedure as Word, with a document-control block."""
    from .documents import is_effective, next_review
    d = model.documents[doc_id]
    meta, body = d["meta"], d["body"]
    blocks = list(_md_blocks(body))
    title = meta.get("title") or doc_id
    if blocks and blocks[0][0] == "h" and blocks[0][1] == 1:
        title = blocks[0][2]
    nr = next_review(model, meta)
    control = [
        ["Document", f"{doc_id} · version {meta.get('version', '—')} · {meta.get('kind', '')}"],
        ["Status", str(meta.get("status") or "—")],
        ["Owner", meta.get("owner") or "—"],
        ["Approved", f"{meta.get('approver') or '—'} on {meta.get('approved_on') or '—'}"],
        ["Next review", str(nr) if nr else "—"],
        ["Requirements", ", ".join(meta.get("satisfies") or []) or "—"],
    ]
    banner = None if is_effective(model, doc_id) else f"DRAFT — {meta.get('status') or 'draft'}, not an approved document"
    first = re.sub(r"^\s*#\s+(.+)$", lambda m: f"# {model.client} — {m.group(1)}", body.lstrip("\n"), count=1, flags=re.M) \
        if blocks and blocks[0][0] == "h" and blocks[0][1] == 1 else f"# {model.client} — {title}\n\n" + body
    return render_markdown(model, first, out_path, f"{model.client} — {title}", safe_stamp(meta), banner, control)


def safe_stamp(meta: dict):
    """Approved: the approval date. Draft: the date it was drafted. Either way, the same text gives the same file."""
    from .documents import safe_date
    drafted = meta.get("drafted") if isinstance(meta.get("drafted"), dict) else {}
    return safe_date(meta.get("approved_on")) or safe_date(drafted.get("on")) or dt.date(2000, 1, 1)
