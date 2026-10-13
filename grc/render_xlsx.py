"""The engagement record as one formatted Excel workbook, for clients and analysts who work in spreadsheets.

A read-only export: the YAML in the repository stays the system of record, and changes go through PRs.
Summary figures are live formulas over the other sheets (COUNTIFS, SUMIFS), so a filtered or sorted copy
keeps adding up. The workbook is reproducible: same record and as-of date, same bytes.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import LEGS, Engagement, to_date

FONT = "Arial"
NAVY = "183B4E"
FILLS = {  # status -> fill colour
    "verified": "C6EFCE", "pass": "C6EFCE", "Ready": "C6EFCE", "approved": "C6EFCE", "closed": "C6EFCE", "done": "C6EFCE",
    "current": "C6EFCE",
    "partial": "FFEB9C", "In progress": "FFEB9C", "in_review": "FFEB9C", "action_planned": "FFEB9C", "implemented": "FFEB9C",
    "claimed": "DDEBF7", "Unverified claim": "DDEBF7", "draft": "DDEBF7",
    "FAILING": "FFC7CE", "EXPIRED": "FFC7CE", "missing": "FFC7CE", "open": "FCE4D6", "overdue": "FFC7CE",
}
THIN = Side(style="thin", color="D9D9D9")


def _date(v):
    try:
        d = to_date(v)
    except (TypeError, ValueError):
        return v
    return d if d else (v or None)


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int],
           status_cols: tuple[int, ...] = (), date_cols: tuple[int, ...] = (), note: str | None = None):
    """A sheet with a frozen, filterable header row, wrapped text and status colouring."""
    ws = wb.create_sheet(title)
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name=FONT, size=9, italic=True, color="666666")
        start = 3
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=start, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r_i, row in enumerate(rows, start=start + 1):
        for c_i, v in enumerate(row, start=1):
            if isinstance(v, (list, tuple)):
                v = ", ".join(str(x) for x in v)
            if c_i in date_cols:
                v = _date(v)
            c = ws.cell(row=r_i, column=c_i, value=v if v not in ("",) else None)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=THIN)
            if isinstance(v, dt.date):
                c.number_format = "yyyy-mm-dd"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    last = start + max(len(rows), 1)
    ref = f"A{start}:{get_column_letter(len(headers))}{last}"
    if rows:
        ws.auto_filter.ref = ref
    for col in status_cols:
        letter = get_column_letter(col)
        rng = f"{letter}{start + 1}:{letter}{last}"
        for value, color in FILLS.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'],
                                                          fill=PatternFill("solid", fgColor=color, bgColor=color)))
    ws.sheet_view.zoomScale = 90
    _print_setup(ws, start)
    return ws, start


def _print_setup(ws, header_row: int | None = None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"


def render_workbook(model: Engagement, path: Path, include_hours: bool = True) -> Path:
    from . import documents, hours, impact, review
    from .rules import leg_effective, metrics, requirement_state
    wb = Workbook()
    wb.remove(wb.active)
    m = metrics(model)
    fw = model.catalog["framework"]
    note = (f"{model.client} · {fw.get('short_name')} · generated from the GRC record as of {model.today}. "
            "Read-only export: change the record through pull requests, not here.")

    # ---------------------------------------------------------- data sheets
    cm_rows = []
    for r in model.requirements:
        c = model.controls.get(r["id"], {})
        excluded = not model.is_applicable(r["id"])
        legs = ["excluded" if excluded else leg_effective(model, (c.get("test") or {}).get(k) or {}) for k in LEGS]
        ev = sorted({str(x) for k in LEGS for x in (((c.get("test") or {}).get(k) or {}).get("evidence") or [])})
        cm_rows.append([r["id"], r["title"], "Clause" if r["kind"] == "clause" else "Annex A",
                        "No" if excluded else "Yes", c.get("owner") or "", *legs, requirement_state(model, r["id"]),
                        ", ".join(c.get("gaps") or []), ", ".join(ev)])
    _sheet(wb, "Control matrix", ["ID", "Requirement", "Type", "Applicable", "Owner", "Documented", "Implemented",
                                  "Evidenced", "Tested", "State", "Gaps", "Evidence"], cm_rows,
           [9, 42, 9, 11, 24, 13, 13, 13, 13, 16, 14, 30], status_cols=(6, 7, 8, 9, 10), note=note)

    soa_rows = [[r["id"], r["title"], "Yes" if model.is_applicable(r["id"]) else "No",
                 model.controls.get(r["id"], {}).get("justification") or "",
                 model.controls.get(r["id"], {}).get("implementation") or "", requirement_state(model, r["id"]),
                 model.controls.get(r["id"], {}).get("owner") or ""]
                for r in model.requirements if r["kind"] == "control"]
    _sheet(wb, "SoA", ["Control", "Title", "Applicable", "Justification", "Implementation", "State", "Owner"], soa_rows,
           [9, 40, 11, 45, 55, 16, 24], status_cols=(6,), note=note)

    gap_rows = []
    for i, g in enumerate(model.gaps, start=4):
        gap_rows.append([g.get("id"), g.get("area"), model.fmt(str(g.get("finding", ""))), ", ".join(map(str, g.get("requirements") or [])),
                         g.get("owner") or "", g.get("due"), g.get("phase") or "", g.get("severity") or "", g.get("status"),
                         f'=IF(AND(I{i}<>"closed",ISNUMBER(F{i}),F{i}<Summary!$B$4),Summary!$B$4-F{i},"")'])
    _sheet(wb, "Gaps", ["ID", "Area", "Finding", "Requirements", "Owner", "Due", "Phase", "Severity", "Status",
                        "Days overdue"], gap_rows, [9, 20, 55, 20, 22, 12, 13, 10, 13, 12],
           status_cols=(9,), date_cols=(6,), note=note)

    supports: dict[str, list[str]] = {}
    for cid, c in model.controls.items():
        for k in LEGS:
            for e in (((c.get("test") or {}).get(k) or {}).get("evidence") or []):
                supports.setdefault(str(e), []).append(f"{cid}·{k[0].upper()}")
    ev_rows = []
    for e in sorted(model.evidence_index.values(), key=lambda x: str(x["id"])):
        if e.get("source") in ("audit", "impact", "document", "review") and e["id"] not in supports:
            continue
        exp = model.evidence_expiry(e)
        status = "FAILING" if model.evidence_is_failing(e) else "EXPIRED" if model.evidence_is_expired(e) else "current"
        ev_rows.append([e["id"], e.get("title"), e.get("type"), e.get("location"), e.get("collected"), exp, status,
                        e.get("approved_by") or "", ", ".join(supports.get(e["id"], [])) or "(unused)"])
    _sheet(wb, "Evidence", ["ID", "Title", "Type", "Location", "Collected", "Expires", "Status", "Approved by", "Supports"],
           ev_rows, [16, 45, 11, 40, 12, 12, 11, 22, 30], status_cols=(7,), date_cols=(5, 6), note=note)

    f_rows = [[f.get("id"), f.get("grade"), f.get("source"), f.get("requirement"), f.get("status"), f.get("owner") or "",
               f.get("due"), f.get("description"), f.get("root_cause") or "", f.get("corrective_action") or ""]
              for f in model.findings if isinstance(f, dict)]
    _sheet(wb, "Findings", ["ID", "Grade", "Source", "Requirement", "Status", "Owner", "Due", "Description", "Root cause",
                            "Corrective action"], f_rows, [9, 8, 15, 11, 15, 22, 12, 50, 35, 35],
           status_cols=(5,), date_cols=(7,), note=note)

    active = impact.active_by_system(model)
    ia_rows = []
    for s in impact.required_systems(model):
        ia = active.get(s["id"]) or {}
        res = [impact.score(x.get("residual_likelihood"), x.get("residual_severity"))
               for x in ia.get("impacts") or [] if isinstance(x, dict)]
        res = [x for x in res if x is not None]
        top = max(res) if res else None
        ia_rows.append([s["id"], s.get("name"), s.get("risk_tier") or "", ia.get("id") or "missing", ia.get("status") or "missing",
                        str(ia.get("decision") or "").replace("_", " "), top, impact.band(model, top) if top else "",
                        impact.next_review(model, ia) if ia else None,
                        "Yes" if ia and impact.is_effective(model, ia) else "No"])
    _sheet(wb, "Impact assessments", ["System", "Name", "Risk tier", "Assessment", "Status", "Decision", "Highest residual",
                                      "Band", "Next review", "Current and approved"], ia_rows,
           [10, 32, 10, 12, 12, 22, 10, 10, 12, 12], status_cols=(4, 5), date_cols=(9,), note=note)

    lib = documents.library(model)
    d_rows = []
    for did in list(lib) + [d for d in documents.docs(model) if d not in lib]:
        d = documents.docs(model).get(did)
        meta = d["meta"] if d else {}
        d_rows.append([did, meta.get("title") or lib.get(did, {}).get("title"), meta.get("kind") or lib.get(did, {}).get("kind"),
                       meta.get("status") or "missing", meta.get("version") or "", meta.get("owner") or "",
                       meta.get("approver") or "", meta.get("approved_on"), documents.next_review(model, meta) if d else None,
                       "Yes" if d and documents.is_effective(model, did) else "No",
                       ", ".join(meta.get("satisfies") or lib.get(did, {}).get("satisfies") or [])])
    _sheet(wb, "Documents", ["ID", "Title", "Type", "Status", "Version", "Owner", "Approver", "Approved", "Next review",
                             "Current and approved", "Requirements"], d_rows, [17, 40, 11, 12, 9, 22, 22, 12, 12, 12, 30],
           status_cols=(4,), date_cols=(8, 9), note=note)

    a_rows = [[a["review"], a.get("id"), a.get("action"), a.get("owner") or "", a.get("due"), a.get("status"),
               a.get("closed_on"), ", ".join(str(x) for x in a.get("evidence") or [])] for a in review.actions(model)]
    _sheet(wb, "Review actions", ["Review", "Action ID", "Action", "Owner", "Due", "Status", "Closed", "Evidence"], a_rows,
           [13, 10, 55, 22, 12, 11, 12, 24], status_cols=(6,), date_cols=(5, 7), note=note)

    if include_hours and hours.tracked(model):
        acts = {a["id"]: a.get("title", a["id"]) for a in hours.method(model)["activities"]}
        h_rows = [[e.get("date"), e.get("person"), e.get("hours"), acts.get(e.get("activity"), e.get("activity")),
                   hours.phase_for(model, e) or "", ", ".join(e.get("refs") or []), e.get("note") or ""]
                  for e in sorted(hours.entries(model), key=lambda e: (str(e.get("date")), str(e.get("person"))))]
        ws, start = _sheet(wb, "Hours", ["Date", "Person", "Hours", "Activity", "Phase", "Refs", "Note"], h_rows,
                           [12, 22, 8, 32, 12, 20, 50], date_cols=(1,), note=note)
        # phase plan vs actual, computed with SUMIFS over the entries
        plan = hours.planned(model)
        col = 9
        ws.cell(row=start, column=col, value="Phase").font = Font(name=FONT, bold=True)
        ws.cell(row=start, column=col + 1, value="Planned").font = Font(name=FONT, bold=True)
        ws.cell(row=start, column=col + 2, value="Used").font = Font(name=FONT, bold=True)
        ws.cell(row=start, column=col + 3, value="Remaining").font = Font(name=FONT, bold=True)
        last = start + max(len(h_rows), 1)
        for i, p in enumerate(hours.phases(model), start=start + 1):
            ws.cell(row=i, column=col, value=str(p["id"])).font = Font(name=FONT)
            pc = ws.cell(row=i, column=col + 1, value=plan.get(str(p["id"]), 0))
            pc.font = Font(name=FONT, color="0000FF")
            ws.cell(row=i, column=col + 2, value=f'=SUMIFS($C${start + 1}:$C${last},$E${start + 1}:$E${last},I{i})').font = Font(name=FONT)
            ws.cell(row=i, column=col + 3, value=f"=J{i}-K{i}").font = Font(name=FONT)
            ws.conditional_formatting.add(f"L{i}", CellIsRule(operator="lessThan", formula=["0"],
                                                              fill=PatternFill("solid", fgColor="FFC7CE", bgColor="FFC7CE")))
        note_row = start + len(hours.phases(model)) + 2
        ws.cell(row=note_row, column=col, value="Planned hours (blue) come from the phase plan in engagement.yaml, or the "
                "playbook's phase split of the budget.").font = Font(name=FONT, size=9, italic=True, color="666666")
        for letter, w in zip("IJKL", (14, 10, 10, 11)):
            ws.column_dimensions[letter].width = w

    # ---------------------------------------------------------- summary (first sheet), formulas over the sheets above
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 60
    ws["A1"] = f"{model.client} — {fw.get('short_name')} readiness workbook"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")
    rows = [("As of", model.today, "The date the record was evaluated; overdue figures count from here."),
            ("Applicable requirements", "=COUNTIFS('Control matrix'!D:D,\"Yes\")", ""),
            ("Requirements ready", "=COUNTIFS('Control matrix'!J:J,\"Ready\")", "Every leg verified on current evidence."),
            ("Share ready", "=IF(B5=0,0,B6/B5)", ""),
            ("Evidence-test legs verified",
             "=COUNTIFS('Control matrix'!F:F,\"verified\")+COUNTIFS('Control matrix'!G:G,\"verified\")"
             "+COUNTIFS('Control matrix'!H:H,\"verified\")+COUNTIFS('Control matrix'!I:I,\"verified\")", "Out of 4 per applicable requirement."),
            ("Share of legs verified", "=IF(B5=0,0,B8/(B5*4))", ""),
            ("Open gaps", "=COUNTA(Gaps!A4:A10000)-COUNTIFS(Gaps!I:I,\"closed\")", ""),
            ("Overdue gaps", "=COUNTIFS(Gaps!J:J,\">0\")", "Not closed and past their due date."),
            ("Open nonconformities (major)", "=COUNTIFS(Findings!B:B,\"major\",Findings!E:E,\"<>closed\",Findings!E:E,\"<>accepted\")", ""),
            ("Open nonconformities (minor)", "=COUNTIFS(Findings!B:B,\"minor\",Findings!E:E,\"<>closed\",Findings!E:E,\"<>accepted\")", ""),
            ("Evidence failing or expired", "=COUNTIFS(Evidence!G:G,\"FAILING\")+COUNTIFS(Evidence!G:G,\"EXPIRED\")", ""),
            ("AI systems with a current, approved impact assessment", "=COUNTIFS('Impact assessments'!J:J,\"Yes\")",
             f"Out of {len(ia_rows)} AI systems in scope."),
            ("Policies and procedures current and approved", "=COUNTIFS(Documents!J:J,\"Yes\")", f"Out of {len(lib)} in the library."),
            ("Management review actions open", "=COUNTIFS('Review actions'!F:F,\"open\")", "")]
    if include_hours and hours.tracked(model):
        budget = float((model.engagement.get("budget") or {}).get("hours_total") or 0)
        rows += [("Hours budget", budget, "From engagement.yaml (budget.hours_total)."),
                 ("Hours used", "=SUM(Hours!C:C)", ""),
                 ("Share of budget used", f"=IF(B{4 + len(rows)}=0,0,B{5 + len(rows)}/B{4 + len(rows)})", "")]
    for i, (label, value, comment) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(name=FONT, size=10)
        c = ws.cell(row=i, column=2, value=value)
        c.font = Font(name=FONT, size=10, bold=True, color="0000FF" if not str(value).startswith("=") else "000000")
        if isinstance(value, dt.date):
            c.number_format = "yyyy-mm-dd"
        if label.startswith("Share"):
            c.number_format = "0.0%"
        ws.cell(row=i, column=3, value=comment or None).font = Font(name=FONT, size=9, color="666666")
    ws.conditional_formatting.add("B11", FormulaRule(formula=["B11>0"], fill=PatternFill("solid", fgColor="FFC7CE", bgColor="FFC7CE")))
    ws.freeze_panes = "A4"
    _print_setup(ws)

    # ---------------------------------------------------------- reproducible file
    stamp = dt.datetime.combine(model.today, dt.time(12, 0))
    wb.properties.creator = wb.properties.lastModifiedBy = "grc-readiness"
    wb.properties.created = wb.properties.modified = stamp
    wb.properties.title = f"{model.client} readiness workbook"
    wb.calculation.fullCalcOnLoad = True   # Excel and LibreOffice compute the formulas on open
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    _pin_modified(path, stamp)        # openpyxl stamps 'modified' with the current time on save
    from .render_docx import _normalize_zip
    _normalize_zip(path, stamp)
    return path


def _pin_modified(path: Path, stamp: dt.datetime) -> None:
    import re
    import zipfile
    tmp = path.with_suffix(".pin")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>", data)
            zout.writestr(item, data)
    tmp.replace(path)
