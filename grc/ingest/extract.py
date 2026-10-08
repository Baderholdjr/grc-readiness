"""Turn client documents into plain text (and rows, for spreadsheets)."""
from __future__ import annotations

import csv
import io
import subprocess
from pathlib import Path

TEXT_EXT = {".md", ".txt", ".yaml", ".yml", ".json"}
SUPPORTED = TEXT_EXT | {".docx", ".pdf", ".xlsx", ".csv"}
MAX_CHARS_PER_FILE = 200_000


def rows(path: Path) -> list[list[str]] | None:
    """Tabular content for CSV/XLSX (first sheet with data); None for other types."""
    ext = path.suffix.lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return [[c.strip() for c in r] for r in csv.reader(fh)]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            data = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
            data = [r for r in data if any(r)]
            if data:
                return data
        return []
    return None


def text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        out = path.read_text(encoding="utf-8", errors="replace")
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text.strip() for c in r.cells))
        out = "\n".join(parts)
    elif ext == ".pdf":
        out = _pdf(path)
    elif ext in (".csv", ".xlsx"):
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            buf = io.StringIO()
            for ws in wb.worksheets:
                buf.write(f"## Sheet: {ws.title}\n")
                for r in ws.iter_rows(values_only=True):
                    if any(v is not None for v in r):
                        buf.write("\t".join("" if v is None else str(v) for v in r) + "\n")
            out = buf.getvalue()
        else:
            out = path.read_text(encoding="utf-8-sig", errors="replace")
    else:
        raise ValueError(f"unsupported file type {ext}")
    return out[:MAX_CHARS_PER_FILE]


def _pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = [f"[page {i + 1}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages)]
        txt = "\n".join(pages)
        if txt.strip():
            return txt
    except ImportError:
        pass
    proc = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True)
    return proc.stdout
