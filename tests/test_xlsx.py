"""Excel workbook export."""
import datetime as dt
import shutil
from pathlib import Path

import yaml
from openpyxl import load_workbook

from grc import hours, package
from grc.cli import main
from grc.model import load
from grc.render_xlsx import render_workbook

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "engagement"
TODAY = dt.date(2026, 10, 8)


def setup(tmp_path, timesheet=True):
    edir = tmp_path / "engagement"
    shutil.copytree(FIXTURE, edir)
    if timesheet:
        (edir / "hours.yaml").write_text("entries: []\n")
        hours.log_entry(load(edir, today=TODAY), 3, "interviews", "ana", on=TODAY - dt.timedelta(days=2))
    return edir


def test_build_writes_workbook_with_all_sheets(tmp_path):
    edir = setup(tmp_path)
    out = tmp_path / "out"
    assert main(["--engagement", str(edir), "--today", TODAY.isoformat(), "build", "--out", str(out)]) == 0
    xlsx = next(out.glob("*-workbook.xlsx"))
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Summary", "Control matrix", "SoA", "Gaps", "Evidence", "Findings", "Impact assessments",
                             "Documents", "Review actions", "Hours"]
    cm = wb["Control matrix"]
    assert [c.value for c in cm[3]][:4] == ["ID", "Requirement", "Type", "Applicable"]
    m = load(edir, today=TODAY)
    assert cm.max_row - 3 == len(m.requirements) and cm.freeze_panes == "B4" and cm.auto_filter.ref
    s = wb["Summary"]
    labels = {s.cell(r, 1).value: s.cell(r, 2).value for r in range(4, s.max_row + 1)}
    assert labels["As of"].date() == TODAY
    assert labels["Requirements ready"].startswith("=COUNTIFS('Control matrix'!J:J")
    assert labels["Hours used"] == "=SUM(Hours!C:C)"
    g = wb["Gaps"]
    assert g["J4"].value.startswith("=IF(AND(I4<>\"closed\"") and isinstance(g["F4"].value, dt.datetime)
    assert wb["Hours"]["K4"].value.startswith("=SUMIFS(")
    assert "workbook.xlsx" in (out / "README.md").read_text()


def test_workbook_is_reproducible_and_package_excludes_hours(tmp_path):
    edir = setup(tmp_path)
    a = render_workbook(load(edir, today=TODAY), tmp_path / "a.xlsx").read_bytes()
    b = render_workbook(load(edir, today=TODAY), tmp_path / "b.xlsx").read_bytes()
    assert a == b
    package.build(load(edir, today=TODAY), tmp_path / "dist")
    root = next(p for p in (tmp_path / "dist").iterdir() if p.is_dir())
    wb = load_workbook(next(root.glob("01-overview/*-workbook.xlsx")))
    assert "Hours" not in wb.sheetnames
    labels = [wb["Summary"].cell(r, 1).value for r in range(4, wb["Summary"].max_row + 1)]
    assert "Hours used" not in labels


def test_workbook_without_timesheet_and_with_odd_values(tmp_path):
    edir = setup(tmp_path, timesheet=False)
    gaps = yaml.safe_load((edir / "gaps.yaml").read_text())
    gaps["gaps"][0]["due"] = "not a date"
    (edir / "gaps.yaml").write_text(yaml.safe_dump(gaps, sort_keys=False))
    wb = load_workbook(render_workbook(load(edir, today=TODAY), tmp_path / "w.xlsx"))
    assert "Hours" not in wb.sheetnames and wb["Gaps"]["F4"].value == "not a date"
