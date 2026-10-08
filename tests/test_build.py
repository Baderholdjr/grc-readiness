import datetime as dt
import shutil
from pathlib import Path

import docx

from grc.cli import main
from grc.issues import issue_body, write_issue_numbers
from grc.model import load

ROOT = Path(__file__).resolve().parent.parent


def test_build_outputs(tmp_path):
    out = tmp_path / "out"
    assert main(["--today", "2026-10-08", "build", "--out", str(out)]) == 0
    for name in ["control-matrix.md", "soa.md", "gap-register.md", "evidence-index.md", "status.json"]:
        assert (out / name).exists()
    brief = docx.Document(str(out / "example-client-iso42001-brief.docx"))
    heads = [p.text for p in brief.paragraphs if p.style.name == "Heading 1"]
    assert heads == ["1. Client Overview", "2. AI Usage", "3. What is already in place", "4. Gap Analysis",
                     "5. Timeline", "6. Needs", "7. ISO Requirements", "8. Engagement Guardrails", "9. Key Risks"]
    text = "\n".join(p.text for p in brief.paragraphs)
    assert "200 total consulting hours" in text
    # closed / not-in-brief gaps stay out of the brief's gap table
    cells = " ".join(c.text for t in brief.tables for r in t.rows for c in r.cells)
    assert "Model registry entries lacked" not in cells
    assert "No AI impact assessments" in cells


def test_issue_body_has_marker():
    m = load(ROOT / "engagement", today=dt.date(2026, 10, 8))
    body = issue_body(m, m.gaps[0])
    assert body.startswith("<!-- grc:GAP-01 -->")
    assert "engagement/controls/5.2.yaml" in body


def test_write_issue_numbers_preserves_comments(tmp_path):
    p = tmp_path / "gaps.yaml"
    shutil.copy(ROOT / "engagement" / "gaps.yaml", p)
    write_issue_numbers(p, {"GAP-01": 7, "GAP-13": 9})
    write_issue_numbers(p, {"GAP-01": 7})  # idempotent
    text = p.read_text()
    assert text.count("issue: 7") == 1 and "issue: 9" in text
    assert text.startswith("# Gap register.")
