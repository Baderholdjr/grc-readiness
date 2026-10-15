"""SOC 2 (AICPA Trust Services Criteria): catalog, categories, wording, document library, mapping to ISO 27001."""
import collections
import datetime as dt
import json
import re

import yaml

from grc import documents, portal
from grc.cli import main
from grc.model import ENGINE, load
from grc.rules import leg_effective, validate

FW = ENGINE / "frameworks" / "soc2"
TODAY = dt.date(2026, 11, 20)
ISO_WORDS = re.compile(r"Annex A|Stage [12]|Statement of Applicability|ISMS|AIMS|certification body|clause 9\.\d|10\.2")


def codes(m):
    return [f.code for f in validate(m)]


def _new(tmp_path, categories="availability,confidentiality", name="engagement-soc2"):
    root = tmp_path / "repo"
    edir = root / name
    args = ["--engagement", str(edir), "new-engagement", "--client", "Acme Health", "--start", "2026-11-02",
            "--stage1", "2027-03-15", "--stage2", "2027-09-15", "--framework", "soc2",
            "--intake", str(root / ("intake" + name[len("engagement"):])), "--confirm-reset"]
    if categories:
        args[-1:-1] = ["--categories", categories]
    assert main(args) == 0
    return edir


def _set(edir, rid, **kw):
    p = edir / "controls" / f"{rid}.yaml"
    d = yaml.safe_load(p.read_text())
    d.update(kw)
    p.write_text(yaml.safe_dump(d, sort_keys=False))


def _set_leg(edir, rid, leg, status, evidence, owner="CTO"):
    p = edir / "controls" / f"{rid}.yaml"
    d = yaml.safe_load(p.read_text())
    d["owner"] = owner
    d["test"][leg] = {"status": status, "evidence": evidence}
    p.write_text(yaml.safe_dump(d, sort_keys=False))


# ------------------------------------------------------------------ catalog and playbook
def test_catalog_shape():
    cat = yaml.safe_load((FW / "catalog.yaml").read_text())
    reqs = cat["requirements"]
    assert collections.Counter(r["category"] for r in reqs) == {
        "security": 33, "availability": 3, "confidentiality": 2, "processing_integrity": 5, "privacy": 18}
    assert {r["kind"] for r in reqs} == {"criterion"} and not cat["kinds"]["criterion"]["mandatory"]
    ids = [r["id"] for r in reqs]
    assert len(set(ids)) == len(ids)
    assert [r["id"] for r in reqs if r["group"] == "CC6"] == [f"CC6.{i}" for i in range(1, 9)]
    assert {g["id"] for g in cat["groups"]} == {r["group"] for r in reqs}
    assert {c["id"] for c in cat["categories"]} == {r["category"] for r in reqs}
    assert [c["id"] for c in cat["categories"] if c.get("required")] == ["security"]
    for r in reqs:
        assert set(r) <= {"id", "kind", "category", "group", "title", "intent", "evidence_examples"}, r["id"]
        assert r["intent"] and r["evidence_examples"] and r["id"].startswith(r["group"]), r["id"]


def test_document_library_covers_every_criterion_once_and_templates_have_their_sections():
    cat = yaml.safe_load((FW / "catalog.yaml").read_text())
    pb = yaml.safe_load((FW / "playbook.yaml").read_text())
    lib = pb["documents"]["library"]
    covered = collections.Counter(r for e in lib for r in e["satisfies"])
    assert set(covered) == {r["id"] for r in cat["requirements"]} and set(covered.values()) == {1}
    cats = {r["id"]: r["category"] for r in cat["requirements"]}
    for e in lib:
        # a document serves one optional category at most, so leaving a category out drops whole documents
        assert len({cats[r] for r in e["satisfies"]}) == 1, e["id"]
        text = (FW / "documents" / e["template"]).read_text()
        heads = re.findall(r"^## (.+)$", text, re.M)
        assert [h for h in heads if h in e["sections"]] == e["sections"], e["id"]
        assert not ISO_WORDS.search(text) and " AI " not in text, e["template"]
    assert pb["impact_assessment"]["required_for_kinds"] == []
    assert {p["id"] for p in pb["phases"]} == {"validate", "governance", "implement", "test", "stage1", "stage2"}


# ------------------------------------------------------------------ an engagement end to end
def test_new_soc2_engagement_validates_builds_and_speaks_soc2(tmp_path):
    edir = _new(tmp_path)
    m = load(edir, today=TODAY)
    assert m.categories_in_scope == {"security", "availability", "confidentiality"}
    assert len(m.controls) == 38 and not m.uses_impact_assessments and m.system_abbrev == "system"
    assert m.engagement["scope"]["categories"] == ["availability", "confidentiality"]
    assert not m.is_applicable("P1.1") and m.is_applicable("A1.1") and m.is_applicable("CC1.1")
    assert codes(m) == ["W140"]
    prog = yaml.safe_load((edir / "audit" / "program.yaml").read_text())["audits"][0]
    assert prog["title"] == "SOC 2 internal audit - full system"
    assert "CC4.1" in (edir / "audit" / "program.yaml").read_text()
    col = (edir / "collectors.yaml").read_text()
    assert '"CC8.1:implemented"' in col and '"7.5:' not in col

    # documents for categories not chosen are not needed
    log = documents.new(m, "all", mode="templates")
    drafted = sorted(p.stem for p in (edir / "documents").glob("DOC-*.md"))
    assert len(drafted) == 16 and "DOC-SOC-PRIVACY" not in drafted and "DOC-SOC-AVAILABILITY" in drafted, log
    assert all("{{" not in (edir / "documents" / f"{d}.md").read_text() for d in drafted)

    out = tmp_path / "out"
    assert main(["--engagement", str(edir), "--today", TODAY.isoformat(), "build", "--out", str(out)]) == 0
    soa = (out / "soa.md").read_text()
    assert soa.startswith("# Acme Health — Criteria Applicability")
    assert "| P1.1 | Privacy notice | No | Privacy is not in the scope of this engagement." in soa
    assert "| A1.1 | Capacity management | Yes |" in soa
    matrix = (out / "control-matrix.md").read_text()
    assert "## CC6 Logical and physical access controls" in matrix and "P1.1" not in matrix
    assert "| CC1.1 | Integrity and ethical values | Security |" in matrix
    for f in ("soa.md", "control-matrix.md", "brief.md", "client-view.md", "internal-audit.md", "README.md"):
        assert not ISO_WORDS.search((out / f).read_text()), f
    cv = (out / "client-view.md").read_text()
    assert "Observation period start" in cv and "Observation period end" in cv
    assert "SOC 2 examination readiness and gap closure" in (out / "brief.md").read_text()

    # Excel: the applicability sheet is named for SOC 2 and lists every criterion
    import openpyxl
    wb = openpyxl.load_workbook(next(out.glob("*-workbook.xlsx")))
    assert "Applicability" in wb.sheetnames and "SoA" not in wb.sheetnames

    # auditor package
    dist = tmp_path / "dist"
    assert main(["--engagement", str(edir), "--today", TODAY.isoformat(), "package", "--out", str(dist)]) == 0
    idx = next(dist.glob("*/00-auditor-index.md")).read_text()
    assert "| Service auditor (CPA firm) | TBD |" in idx and "| Observation period start / Observation period end |" in idx
    assert "| System scope |" in idx and "| P1.1 |" not in idx and not ISO_WORDS.search(idx)


def test_site_tab_uses_soc2_terms_and_glossary(tmp_path):
    edir = _new(tmp_path)
    d = portal.data(load(edir, today=TODAY))
    f = d["frameworks"][0]
    assert f["name"] == "SOC 2" and f["terms"]["stage1"] == "Observation period start"
    assert [a["id"] for a in f["areas"]] == ["CC1", "CC2", "CC3", "CC4", "CC5", "CC6", "CC7", "CC8", "CC9", "A1", "C1"]
    assert f["areas"][0]["requirements"][0]["kind_label"] == "Trust Services Criterion"
    assert list(dict.fromkeys(a["section"] for a in f["areas"])) == ["Security (common criteria)", "Availability", "Confidentiality"]
    assert not f["areas"][0]["section_note"] and "optional" in f["areas"][-1]["section_note"]
    assert [s["name"] for s in f["path"]["stages"]] == ["Observation period start", "Observation period end"]
    assert not [u for u in f["path"]["upcoming"] if u["text"].startswith("Observation period")]
    assert "Observation period" in [g["term"] for g in f["glossary"]]
    assert f["applicable"] == 38 and not f["impact_assessments"] and f["terms"]["review_ref"] == "CC1.2"
    js = (ENGINE / "portal" / "site" / "app.js").read_text()
    assert "Annex A control" not in js and "f.glossary" in js and 'refs("review_ref")' in js


# ------------------------------------------------------------------ categories and applicability
def test_categories_are_checked_and_can_be_added_later(tmp_path):
    root = tmp_path / "repo"
    assert main(["--engagement", str(root / "engagement-soc2"), "new-engagement", "--client", "A", "--start", "2026-11-02",
                 "--stage1", "2027-03-15", "--stage2", "2027-09-15", "--framework", "soc2", "--categories", "privcy",
                 "--intake", str(root / "intake-soc2"), "--confirm-reset"]) == 2
    edir = _new(tmp_path, categories=None)
    m = load(edir, today=TODAY)
    assert len(m.controls) == 33 and m.categories_in_scope == {"security"}
    # choosing a category later: CI asks for the missing assessment files until `grc init` adds them
    p = edir / "engagement.yaml"
    e = yaml.safe_load(p.read_text())
    e["scope"] = {"categories": ["privacy", "availabilty"]}
    p.write_text(yaml.safe_dump(e, sort_keys=False))
    fs = validate(load(edir, today=TODAY))
    assert [f.code for f in fs].count("E002") == 18
    assert any(f.code == "E016" and "availabilty" in f.message for f in fs)
    e["scope"] = {"categories": ["privacy"]}
    p.write_text(yaml.safe_dump(e, sort_keys=False))
    assert main(["--engagement", str(edir), "init"]) == 0
    assert "E002" not in codes(load(edir, today=TODAY)) and len(load(edir).controls) == 51
    # ISO frameworks have no categories to choose
    iso = tmp_path / "iso" / "engagement"
    assert main(["--engagement", str(iso), "new-engagement", "--client", "A", "--start", "2026-11-02", "--stage1",
                 "2027-03-15", "--stage2", "2027-05-10", "--framework", "iso27001", "--categories", "privacy",
                 "--intake", str(tmp_path / "iso" / "intake"), "--confirm-reset"]) == 2


def test_a_criterion_marked_not_applicable_needs_a_justification(tmp_path):
    edir = _new(tmp_path)
    _set(edir, "CC6.4", applicable=False)
    e4 = [f for f in validate(load(edir, today=TODAY)) if f.code == "E004"]
    assert len(e4) == 1 and e4[0].message.startswith("Trust Services Criterion CC6.4 is excluded without a justification")
    assert "Criteria Applicability" in e4[0].message
    _set(edir, "CC6.4", applicable=False, justification="No offices; production runs at the cloud host (carved out).")
    m = load(edir, today=TODAY)
    assert "E004" not in codes(m) and not m.is_applicable("CC6.4")
    assert m.categories_in_scope == {"security", "availability", "confidentiality"}


def test_iso_wording_is_unchanged(tmp_path):
    root = tmp_path / "repo"
    edir = root / "engagement"
    assert main(["--engagement", str(edir), "new-engagement", "--client", "A", "--start", "2026-11-02", "--stage1",
                 "2027-03-15", "--stage2", "2027-05-10", "--framework", "iso27001", "--intake", str(root / "intake"),
                 "--confirm-reset"]) == 0
    out = tmp_path / "out"
    assert main(["--engagement", str(edir), "--today", TODAY.isoformat(), "build", "--out", str(out), "--no-docx"]) == 0
    assert (out / "soa.csv").read_text().startswith("Control,")
    _set(edir, "4.1", applicable=False)
    _set(edir, "A.5.7", applicable=False)
    msgs = [f.message for f in validate(load(edir, today=TODAY)) if f.code == "E004"]
    assert "Clause 4.1 is mandatory and cannot be excluded." in msgs
    assert "Annex A control A.5.7 is excluded without a justification (needed for the Statement of Applicability)." in msgs
    m = load(edir, today=TODAY)
    assert m.term("stage1") == "Stage 1 audit" and m.kind_label(m.requirement_index["A.5.7"]) == "Annex A"
    from grc.render_md import applicability_rows
    assert len(applicability_rows(m)) == 93
    d = portal.data(m)["frameworks"][0]
    assert [s["name"] for s in d["path"]["stages"]] == ["Stage 1 audit", "Stage 2 audit"]
    assert "Stage 1 and Stage 2 audits" in [g["term"] for g in d["glossary"]]
    assert list(dict.fromkeys(a["section"] for a in d["areas"])) == ["Management system requirements", "Controls"]


# ------------------------------------------------------------------ mapping to ISO 27001
def test_mapping_ids_exist_and_evidence_is_shared_through_a_reviewed_equivalent(tmp_path):
    src = yaml.safe_load((ENGINE / "mappings" / "iso27001__soc2.yaml").read_text())
    iso = {r["id"] for r in yaml.safe_load((ENGINE / "frameworks" / "iso27001" / "catalog.yaml").read_text())["requirements"]}
    soc = {r["id"] for r in yaml.safe_load((FW / "catalog.yaml").read_text())["requirements"]}
    assert src["frameworks"] == ["iso27001", "soc2"]
    assert all(m["iso27001"] in iso and m["soc2"] in soc and not m["reviewed_by"] for m in src["mappings"])
    assert ("A.8.32", "CC8.1") in {(m["iso27001"], m["soc2"]) for m in src["mappings"] if m["relationship"] == "equivalent"}

    esoc = _new(tmp_path)
    root = esoc.parent
    e27 = root / "engagement"
    assert main(["--engagement", str(e27), "new-engagement", "--client", "Acme Health", "--start", "2026-11-02",
                 "--stage1", "2027-03-15", "--stage2", "2027-05-10", "--framework", "iso27001",
                 "--intake", str(root / "intake"), "--confirm-reset"]) == 0
    (e27 / "evidence" / "index.yaml").write_text(yaml.safe_dump({"evidence": [
        {"id": "EV-001", "title": "Branch protection and PR review settings", "type": "config", "location": "github://org",
         "collected": "2026-11-01", "approved_by": "CTO", "approved_on": "2026-11-01", "valid_for_days": 365}]}))
    p = e27 / "controls" / "A.8.32.yaml"
    d = yaml.safe_load(p.read_text())
    d["owner"] = "CTO"
    d["test"]["implemented"] = {"status": "verified", "evidence": ["EV-001"]}
    p.write_text(yaml.safe_dump(d, sort_keys=False))
    _set_leg(esoc, "CC8.1", "implemented", "verified", ["iso27001:EV-001"])
    assert "E080" in codes(load(esoc, today=TODAY))                 # only proposed so far
    for m in src["mappings"]:
        if (m["iso27001"], m["soc2"]) == ("A.8.32", "CC8.1"):
            m.update(reviewed_by="Bernard Aderhold", reviewed_on="2026-11-03")
    (root / "mappings").mkdir()
    (root / "mappings" / "iso27001__soc2.yaml").write_text(yaml.safe_dump(src, sort_keys=False))
    msoc = load(esoc, today=TODAY)
    assert "E080" not in codes(msoc)
    assert leg_effective(msoc, msoc.controls["CC8.1"]["test"]["implemented"]) == "verified"
    # the site shows the two records as two tabs
    site = tmp_path / "site"
    assert main(["--engagement", str(e27), "--today", TODAY.isoformat(), "portal", "--out", str(site), "--repo", "a/b",
                 "--also", str(esoc)]) == 0
    data = json.loads((site / "public" / "portal.json").read_text())
    assert [f["key"] for f in data["frameworks"]] == ["iso27001", "soc2"]
    # the cross-framework report leaves out criteria in categories not chosen (privacy here)
    out = tmp_path / "outsoc"
    assert main(["--engagement", str(esoc), "--today", TODAY.isoformat(), "build", "--out", str(out), "--no-docx"]) == 0
    xf = (out / "cross-framework.md").read_text()
    assert "| CC8.1 |" in xf and "P4.3" not in xf and "Bernard Aderhold" in xf


# ------------------------------------------------------------------ fixes from review
def test_review_pack_fills_every_soc2_input(tmp_path):
    from grc import review
    edir = _new(tmp_path)
    m = load(edir, today=TODAY)
    rid = next(iter(review.reviews(m)))
    text = review.pack_markdown(m, rid)
    for title in ("Control performance", "Changes to the services", "Results of the risk assessment",
                  "Subservice organizations", "Security incidents"):
        sec = text.split(title, 1)[1].split("\n## ", 1)[0]
        assert "No generated data" not in sec, title
    assert "DOC-SCOPE" not in text and "DOC-OBJECTIVES" not in text and "DOC-SOC-RISK" in text
    # SOC 2 doesn't require a management review before the observation period
    assert "W151" not in codes(load(edir, today=dt.date(2027, 3, 1)))


def test_timesheet_refs_recognise_soc2_ids():
    from grc.hours import NOTE_REF, REF
    assert REF.findall("CC6.1;A1.2;PI1.3;P6.1;C1.2;A.8.32;9.2") == ["CC6.1", "A1.2", "PI1.3", "P6.1", "C1.2", "A.8.32", "9.2"]
    assert NOTE_REF.findall("access review CC6.3 took 2.5 hours") == ["CC6.3"]


def test_gaps_and_mappings_outside_the_chosen_categories_are_not_chased(tmp_path):
    from grc import client_view, issues
    from grc.rules import metrics
    edir = _new(tmp_path, categories="availability,privacy")
    (edir / "gaps.yaml").write_text(yaml.safe_dump({"gaps": [
        {"id": "GAP-01", "area": "Privacy notice", "finding": "No notice.", "requirements": ["P1.1"], "owner": "CTO",
         "due": "2026-12-01", "phase": "governance", "severity": "low", "status": "open"}]}))
    m = load(edir, today=TODAY)
    assert "W112" not in codes(m) and metrics(m)["gaps_open"] == 1
    p = edir / "engagement.yaml"
    e = yaml.safe_load(p.read_text())
    e["scope"] = {"categories": ["availability"]}
    p.write_text(yaml.safe_dump(e, sort_keys=False))
    m = load(edir, today=TODAY)
    assert "W112" in codes(m) and metrics(m)["gaps_open"] == 0
    assert not [i for i in client_view.items(m) if i["kind"] == "Gap"]
    assert not next(i for i in issues._items(m) if i["id"] == "GAP-01")["open"]
    assert metrics(m)["requirements_total"] == 36
    # applicability CSV keeps a framework-appropriate first column
    out = tmp_path / "out"
    assert main(["--engagement", str(edir), "--today", TODAY.isoformat(), "build", "--out", str(out), "--no-docx"]) == 0
    assert (out / "soa.csv").read_text().startswith("Criterion,")
